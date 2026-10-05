"""
appendix.py — Detailed per-slice workbooks, one per evaluation (thesis appendices).

Purpose : Give every Chapter 4 evaluation an Excel file that lists the result of
          every sampled slice, so each table value can be cross-referenced to the
          slices behind it.
Function : write_all(df, timing, R, d) writes outputs/evaluation/appendix/
            Appendix_A_Sampling.xlsx                 (Tables 4.1–4.2)
            Appendix_B_E1_Seed_Selection.xlsx        (Objective 1; Tables 4.9–4.11)
            Appendix_C_E5_Failure_Attribution.xlsx   (Table 4.12)
            Appendix_D_E2_Undersegmentation.xlsx     (Objective 2; Tables 4.13–4.14)
            Appendix_E_E3_Boundary_Leakage.xlsx      (Objective 3; Tables 4.15–4.17)
            Appendix_F_E4_Overall_Delineation.xlsx   (Tables 4.4–4.8, 4.18–4.20)
            Appendix_G_E6_Processing_Time.xlsx       (Tables 4.21–4.22)
            Appendix_H_Raw_Results.xlsx              (every run, every column)
          Each workbook opens with a Notes sheet (purpose, equations, column
          dictionary), then per-slice sheets, then the Chapter 4 tables it supports.
Notes   : Columns whose header ends in "(check)" are live Excel formulas that
          recompute the metric from the pixel counts in the same row, so a reader
          can verify the stored value without any software but Excel.
"""
import os
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from experiments.chapter4 import tables as T
from experiments.evaluate import CONFIGS

FONT = "Times New Roman"
THIN = Side(style="thin", color="000000")
F4, F3, F2, INT, PCT = "0.0000", "0.000", "0.00", "#,##0", "0.0%"
CLS = {"glioma": "Glioma", "meningioma": "Meningioma", "pituitary": "Pituitary"}
PLN = {"axial": "Axial", "coronal": "Coronal", "sagittal": "Sagittal"}


def plain(s):
    """Renderer markup -> plain text for Excel."""
    s = str(s)
    s = s.replace("**", "").replace("*", "")
    s = re.sub(r"\^\{2\}", "²", s)
    s = re.sub(r"_\{([^}]*)\}", lambda m: m.group(1) if len(m.group(1)) == 1 else f"_{m.group(1)}", s)
    s = re.sub(r"\^\{([^}]*)\}", r"^\1", s)
    return s


# ─── Sheet writers ───────────────────────────────────────────────────────────
def _title(ws, label, title, desc=None):
    ws["A1"] = label
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    ws["A2"] = title
    ws["A2"].font = Font(name=FONT, italic=True, size=12)
    if desc:
        ws["A3"] = desc
        ws["A3"].font = Font(name=FONT, size=10)
        ws["A3"].alignment = Alignment(wrap_text=False)


def data_sheet(wb, name, label, title, df, fmt=None, checks=None, desc=None, widths=None):
    """
    One row per slice. fmt maps header -> number format; checks is a list of
    (header, template) where {Header Name} is replaced by that column's cell.
    """
    ws = wb.create_sheet(name)
    _title(ws, label, title, desc)
    df = df.reset_index(drop=True)
    headers = list(df.columns) + [h for h, _ in (checks or [])]
    hr = 5
    col = {h: get_column_letter(i + 1) for i, h in enumerate(headers)}
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=hr, column=j, value=h)
        c.font = Font(name=FONT, bold=True, size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=THIN, bottom=THIN)
    vals = df.to_numpy(dtype=object)
    for i in range(len(df)):
        r = hr + 1 + i
        for j, v in enumerate(vals[i], 1):
            if isinstance(v, (float, np.floating)) and np.isnan(v):
                v = None
            elif isinstance(v, np.generic):
                v = v.item()
            ws.cell(row=r, column=j, value=v)
        for k, (h, tpl) in enumerate(checks or []):
            formula = re.sub(r"\{([^}]+)\}", lambda m: f"{col[m.group(1)]}{r}", tpl)
            ws.cell(row=r, column=len(df.columns) + 1 + k, value=f'=IFERROR({formula[1:]},"")')
    last = hr + len(df)
    for j, h in enumerate(headers, 1):
        f = (fmt or {}).get(h) or (F4 if h.endswith("(check)") else None)
        L = get_column_letter(j)
        if f:
            for r in range(hr + 1, last + 1):
                ws[f"{L}{r}"].number_format = f
        ws.column_dimensions[L].width = (widths or {}).get(h, max(11, min(34, len(h) * 0.9 + 2)))
    for j in range(1, len(headers) + 1):
        ws.cell(row=last, column=j).border = Border(bottom=THIN)
    ws.row_dimensions[hr].height = 42
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{last}"
    return ws


def notes_sheet(wb, label, title, items):
    ws = wb.active
    ws.title = "Notes"
    _title(ws, label, title)
    r = 4
    for k, v in items:
        a = ws.cell(row=r, column=1, value=k)
        a.font = Font(name=FONT, bold=True, size=11)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b = ws.cell(row=r, column=2, value=v)
        b.font = Font(name=FONT, size=11)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120


def table_sheet(wb, name, tabs):
    """Chapter 4 tables in APA layout (number bold, title italic, rules above/below the header and at the end)."""
    ws = wb.create_sheet(name)
    r = 1
    for t in tabs:
        ws.cell(row=r, column=1, value=f"Table {t['num']}").font = Font(name=FONT, bold=True, size=12)
        ws.cell(row=r + 1, column=1, value=plain(t["title"])).font = Font(name=FONT, italic=True, size=12)
        r += 2
        hdrs = t["header"] if t["header"] and isinstance(t["header"][0], list) else [t["header"]]
        ncol = len(t["widths"])
        for hi, h in enumerate(hdrs):
            j = 1
            for c in h:
                o = c if isinstance(c, dict) else {"text": c}
                cell = ws.cell(row=r, column=j, value=plain(o["text"]))
                cell.font = Font(name=FONT, bold=True, size=10)
                cell.alignment = Alignment(horizontal="center" if j > 1 else "left", wrap_text=True, vertical="center")
                span = o.get("span", 1)
                if span > 1:
                    ws.merge_cells(start_row=r, start_column=j, end_row=r, end_column=j + span - 1)
                j += span
            for jj in range(1, ncol + 1):
                ws.cell(row=r, column=jj).border = Border(top=THIN if hi == 0 else None,
                                                          bottom=THIN if hi == len(hdrs) - 1 else None)
            r += 1
        for row in t["rows"]:
            if isinstance(row, dict):
                c = ws.cell(row=r, column=1, value=plain(row["group"]))
                c.font = Font(name=FONT, italic=True, size=10)
            else:
                for j, v in enumerate(row, 1):
                    c = ws.cell(row=r, column=j, value=plain(v["text"] if isinstance(v, dict) else v))
                    c.font = Font(name=FONT, size=10)
                    c.alignment = Alignment(horizontal="left" if j == 1 else "center", wrap_text=True)
            r += 1
        for jj in range(1, ncol + 1):
            ws.cell(row=r - 1, column=jj).border = Border(bottom=THIN)
        if t.get("note"):
            c = ws.cell(row=r, column=1, value="Note. " + plain(t["note"]))
            c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
            ws.row_dimensions[r].height = max(30, 13 * (len(t["note"]) // (ncol * 14) + 1))
        r += 3
    ws.column_dimensions["A"].width = 34
    for j in range(2, 12):
        ws.column_dimensions[get_column_letter(j)].width = 17


def save(wb, d, name):
    os.makedirs(os.path.join(d, "appendix"), exist_ok=True)
    p = os.path.join(d, "appendix", name)
    wb.save(p)
    return p


# ─── Helpers ─────────────────────────────────────────────────────────────────
def base_cols(df):
    a = df[df.config == "A_ESRG"].set_index("file")
    out = pd.DataFrame({"Slice (file)": a.index, "Split": a.split.values,
                        "Tumor class": a.tumor.map(CLS).values, "Plane": a.plane.map(PLN).values,
                        "Ground-truth area |G| (px)": a.gt_area.values})
    return out.sort_values("Slice (file)").set_index("Slice (file)", drop=False)


def block(df, cfg, cols, prefix):
    """Selected columns of one configuration, renamed "<prefix>: <name>"."""
    s = df[df.config == cfg].set_index("file")
    return pd.DataFrame({f"{prefix}: {name}": s[c] for c, name in cols})


def bool01(s):
    return s.map(lambda v: None if v is None or (isinstance(v, float) and np.isnan(v)) else int(bool(v)))


OVERLAP = [("tp", "TP (px)"), ("fp", "FP (px)"), ("fn", "FN (px)"), ("pred_area", "|M| (px)"),
           ("dsc", "DSC"), ("iou", "IoU"), ("precision", "Precision"), ("recall", "Recall"),
           ("hd95_wc", "HD95 (px)"), ("assd_wc", "ASSD (px)"), ("area_ratio", "Area ratio |M|/|G|"),
           ("leaked", "Leaked (1 = yes)"), ("success", "Success (1 = yes)")]


def fmt_for(columns):
    out = {}
    for c in columns:
        if any(k in c for k in ("(px)", "TP", "FP", "FN")) and "HD95" not in c and "ASSD" not in c and "distance" not in c:
            out[c] = INT
        elif any(k in c for k in ("HD95", "ASSD", "distance", "depth")):
            out[c] = F2
        elif any(k in c for k in ("DSC", "IoU", "Precision", "Recall", "share", "ratio", "Missed", "SD", "range", "mean")):
            out[c] = F4
    return out


def overlap_checks(prefix):
    return [(f"{prefix}: DSC (check)", f"=2*{{{prefix}: TP (px)}}/(2*{{{prefix}: TP (px)}}+{{{prefix}: FP (px)}}+{{{prefix}: FN (px)}})"),
            (f"{prefix}: Recall (check)", f"={{{prefix}: TP (px)}}/({{{prefix}: TP (px)}}+{{{prefix}: FN (px)}})"),
            (f"{prefix}: Precision (check)", f"={{{prefix}: TP (px)}}/({{{prefix}: TP (px)}}+{{{prefix}: FP (px)}})")]


def metric_notes():
    return [
        ("DSC (Eq. 3.32)", "DSC = 2TP / (2TP + FP + FN) = 2|M ∩ G| / (|M| + |G|)"),
        ("IoU (Eq. 3.33)", "IoU = TP / (TP + FP + FN)"),
        ("Precision (Eq. 3.34)", "Precision = TP / (TP + FP); undefined (blank) when nothing was predicted"),
        ("Recall (Eq. 3.35)", "Recall = TP / (TP + FN); the missed share is 1 − Recall = FN / |G|"),
        ("HD95 (Eq. 3.37)", "95th percentile of the pooled boundary-to-boundary distances (pixels); "
                            "an empty prediction is given the image diagonal (worst case)"),
        ("ASSD (Eq. 4.10)", "Mean of the pooled boundary-to-boundary distances (pixels); worst case as for HD95"),
        ("Leaked (Eq. 3.39)", "1 when |M| > 2|G| (λ = 2)"),
        ("Success (Eq. 4.11)", "1 when DSC ≥ 0.70"),
        ("TP, FP, FN", "Pixel counts: TP = |M ∩ G|, FP = |M \\ G|, FN = |G \\ M| (M = predicted mask, G = ground truth)"),
        ("(check) columns", "Live Excel formulas recomputing the metric from TP, FP, FN in the same row; they "
                            "equal the stored value to rounding precision"),
    ]


# ─── Workbooks ───────────────────────────────────────────────────────────────
def wb_sampling(R, tabs, d):
    wb = Workbook()
    notes_sheet(wb, "Appendix A", "Sample-Size Determination and Stratified Sample", [
        ("Population", f"BRISC 2025 segmentation task, training + test splits: {R['meta']['N_population']:,} slices, "
                       "every one with a ground-truth tumor mask"),
        ("Strata", "9 strata = 3 tumor classes × 3 imaging planes, read from each filename"),
        ("Cochran (Eq. 4.1)", "n0 = z² p (1 − p) / e²,  z = 1.959964, p = 0.5, e = 0.05"),
        ("Finite population correction (Eq. 4.2)", "n = n0 / (1 + (n0 − 1) / N)"),
        ("Equal allocation (Eq. 4.3)", "n_h = ceil( max_c n_c / 3 ) for every stratum h"),
        ("Achieved margin (Eq. 4.4)", "e = z sqrt( p(1 − p)/n × (N − n)/(N − 1) )"),
        ("Selection", "Simple random sampling without replacement within each stratum; the stratum’s slices "
                      "are sorted by filename, shuffled with random.Random(2026), and the first n_h are taken"),
        ("Sheets", "Sample_size (with live formulas), Strata, Sample (the 954 slices), Sampling_frame (all slices "
                   "with their random draw order), Tables (Tables 4.1–4.2)"),
        ("Removal", "Every image/mask pair not in the sample was deleted from segmentation_task/ after the draw "
                    "(experiments/sampling.py --prune); the frame below is the record of the population"),
    ])
    ss = pd.DataFrame(R["sample_size"]).rename(columns={
        "level": "Population", "N": "N (slices)", "n0": "n0 (stored)", "n_exact": "n (stored)", "n_min": "n_min",
        "n_drawn": "n drawn", "margin_achieved": "Achieved e (stored)"})
    ss["z"], ss["p"], ss["e"] = 1.959964, 0.5, 0.05
    data_sheet(wb, "Sample_size", "Appendix A", "Cochran Sample Size per Tumor Class", ss,
               fmt={"n0 (stored)": F4, "n (stored)": F4, "Achieved e (stored)": F4, "N (slices)": INT},
               checks=[("n0 (check)", "={z}^2*{p}*(1-{p})/{e}^2"),
                       ("n (check)", "={n0 (check)}/(1+({n0 (check)}-1)/{N (slices)})"),
                       ("Achieved e (check)", "={z}*SQRT({p}*(1-{p})/{n drawn}*({N (slices)}-{n drawn})/({N (slices)}-1))")])
    st = pd.DataFrame(R["strata"]).rename(columns={"tumor": "Tumor class", "plane": "Plane", "N_h": "N_h",
                                                     "N_h_train": "N_h from train", "N_h_test": "N_h from test",
                                                     "n_h": "n_h", "n_h_train": "n_h from train",
                                                     "n_h_test": "n_h from test", "f_h": "f_h = n_h/N_h"})
    data_sheet(wb, "Strata", "Appendix A", "Population and Sample per Stratum", st, fmt={"f_h = n_h/N_h": PCT},
               checks=[("W_h = N_h/N (check)", "={N_h}/SUM(C:C)")])
    smp = pd.read_csv(os.path.join(d, "sample.csv"))
    smp = smp.rename(columns={"file": "Slice (file)", "split": "Split", "index": "Index", "tumor": "Tumor class",
                              "plane": "Plane", "stratum": "Stratum", "draw_order": "Random draw order",
                              "selected": "Selected", "image": "Image path", "mask": "Mask path"})
    data_sheet(wb, "Sample", "Appendix A", "The 954 Sampled Slices", smp)
    fr = pd.read_csv(os.path.join(d, "frame.csv")).rename(columns={
        "file": "Slice (file)", "split": "Split", "index": "Index", "tumor": "Tumor class", "plane": "Plane",
        "stratum": "Stratum", "draw_order": "Random draw order", "selected": "Selected", "image": "Image path",
        "mask": "Mask path"})
    data_sheet(wb, "Sampling_frame", "Appendix A", "Every Slice of the Population With Its Random Draw Order", fr,
               desc="A slice is selected when its draw order within its stratum is ≤ n_h.")
    table_sheet(wb, "Tables", [tabs["sample_size"], tabs["strata"]])
    return save(wb, d, "Appendix_A_Sampling.xlsx")


def wb_e1(df, R, tabs, d):
    wb = Workbook()
    notes_sheet(wb, "Appendix B", "Experiment E1 — Objective 1: Automated Seed Selection", [
        ("Objective", "To provide an automatic seed selection option through Multi-Level Otsu Thresholding"),
        ("Configuration", "A_ESRG (automatic seed); O_SRG_1–5 and O_ESRG_1–5 (five simulated operator clicks)"),
        ("Seed Hit (Eq. 3.38)", "Hit = 1 when |S \\ G| = 0, i.e. every seed-core pixel lies inside the tumor"),
        ("Seed Hit Rate", "SHR = hits / N_S, N_S = slices that received a seed; Wilson 95% CI (Eq. 4.14)"),
        ("Inside share", "|S ∩ G| / |S|"),
        ("Localization distance (Eq. 4.7)", "Euclidean distance (px) from the centroid of S to the nearest tumor pixel; 0 inside"),
        ("Operator click (Eq. 4.5)", "Random tumor pixel whose distance to the tumor boundary exceeds the 3-px click radius; "
                                     "5 distinct clicks per slice, seeded by the filename (CRC32)"),
        ("Within-slice SD (Eq. 4.8)", "SD_i = sqrt( Σ_k (DSC_ik − mean_i)² / (K − 1) ), K = 5"),
        ("Inconsistency (Eq. 4.9)", "1 when 0 < number of successful clicks < 5"),
        ("Determinism", "Automatic seeding run twice (parallel run and sequential timing pass); identical = same DSC and |M|"),
        ("Sheets", "Seed_selection, Operator_variability, Determinism, Tables"),
    ] + metric_notes())
    b = base_cols(df)
    a = df[df.config == "A_ESRG"].set_index("file")
    s = df[df.config == "A_SRG"].set_index("file")
    seed = pd.DataFrame({
        "Seed core |S| (px)": a.seed_area, "|S ∩ G| (px)": a.seed_in_gt_px,
        "|S \\ G| (px)": a.seed_area - a.seed_in_gt_px, "No candidate (1 = yes)": (a.seed_area == 0).astype(int),
        "Seed hit (1 = yes)": bool01(a.seed_hit), "Inside share |S∩G|/|S|": a.seed_in_gt_frac,
        "Localization distance (px)": a.seed_distance, "ESRG auto: DSC": a.dsc, "ESRG auto: Success (1 = yes)": bool01(a.success),
        "SRG auto: DSC": s.dsc, "Status": a.status})
    out = b.join(seed).reset_index(drop=True)
    data_sheet(wb, "Seed_selection", "Appendix B", "Automated Seed per Slice (Configuration A_ESRG)", out,
               fmt=fmt_for(out.columns),
               checks=[("Seed hit (check)", '=IF({Seed core |S| (px)}=0,"",IF({|S \\ G| (px)}=0,1,0))'),
                       ("Inside share (check)", "={|S ∩ G| (px)}/{Seed core |S| (px)}")])
    per = {k: pd.DataFrame(v).set_index("file") for k, v in R["e1"]["operator_per"].items()}
    cols = {}
    for k in range(1, 6):
        o = df[df.config == f"O_SRG_{k}"].set_index("file")
        cols[f"Click {k}: row"] = o.click_row
        cols[f"Click {k}: col"] = o.click_col
        cols[f"Click {k}: depth (px)"] = o.click_depth
    for alg in ("SRG", "ESRG"):
        for k in range(1, 6):
            cols[f"{alg} click {k}: DSC"] = df[df.config == f"O_{alg}_{k}"].set_index("file").dsc
        p = per[alg.lower()]
        cols[f"{alg}: mean DSC"] = p["mean"]
        cols[f"{alg}: SD of DSC"] = p["sd"]
        cols[f"{alg}: range of DSC"] = p["range"]
        cols[f"{alg}: successful clicks"] = p["n_success"]
        cols[f"{alg}: inconsistent (1 = yes)"] = p["inconsistent"].astype(int)
    cols["ESRG automatic: DSC"] = a.dsc
    ov = b.join(pd.DataFrame(cols)).reset_index(drop=True)
    fm = fmt_for(ov.columns)
    for c in ov.columns:
        if c.endswith(": row") or c.endswith(": col"):
            fm[c] = INT
    checks = []
    for alg in ("SRG", "ESRG"):
        rng = ",".join("{" + f"{alg} click {k}: DSC" + "}" for k in range(1, 6))
        checks.append((f"{alg}: SD of DSC (check)", f"=STDEV({rng})"))
    data_sheet(wb, "Operator_variability", "Appendix B",
               "DSC From Five Simulated Operator Clicks per Slice (O_SRG_1–5, O_ESRG_1–5)", ov, fmt=fm, checks=checks)
    det = R["e1"].get("determinism")
    if det is not None:
        tp = os.path.join(d, "timing_sequential.csv")
        t = pd.read_csv(tp)
        rows = []
        for cfg in ("A_SRG", "A_ESRG"):
            r1 = df[df.config == cfg].set_index("file")
            r2 = t[t.config == cfg].set_index("file")
            idx = r1.index.intersection(r2.index)
            x = pd.DataFrame({"Slice (file)": idx, "Configuration": cfg,
                              "Run 1 (parallel): DSC": r1.loc[idx, "dsc"].values,
                              "Run 2 (sequential): DSC": pd.to_numeric(r2.loc[idx, "dsc"]).values,
                              "Run 1: |M| (px)": r1.loc[idx, "pred_area"].values,
                              "Run 2: |M| (px)": pd.to_numeric(r2.loc[idx, "pred_area"]).values})
            rows.append(x)
        dd = pd.concat(rows).sort_values(["Configuration", "Slice (file)"])
        data_sheet(wb, "Determinism", "Appendix B", "Automatic Seeding Run Twice on Every Slice", dd,
                   fmt=fmt_for(dd.columns),
                   checks=[("Identical (check)", "=IF(AND(ABS({Run 1 (parallel): DSC}-{Run 2 (sequential): DSC})<1E-12,"
                                                 "{Run 1: |M| (px)}={Run 2: |M| (px)}),1,0)")])
    table_sheet(wb, "Tables", [tabs["seed_hit"], tabs["hit_effect"], tabs["operator"]])
    return save(wb, d, "Appendix_B_E1_Seed_Selection.xlsx")


def wb_e5(df, R, tabs, d):
    wb = Workbook()
    notes_sheet(wb, "Appendix C", "Experiment E5 — Failure Attribution of the Enhanced Algorithm (Automatic Seeding)", [
        ("Configuration", "A_ESRG"),
        ("Rule", "Each slice is assigned to the first failed gate, in order: X, A, B, C, G, F, D, E (Section 3.4.4)"),
    ] + [(f"Bucket {k}", v) for k, v in R["e5_labels"].items() if k != "?"])
    b = base_cols(df)
    a = df[df.config == "A_ESRG"].set_index("file")
    out = b.join(pd.DataFrame({"Bucket": a.bucket, "Bucket meaning": a.bucket.map(R["e5_labels"]),
                               "Measured value that set the bucket": a.bucket_reason, "DSC": a.dsc,
                               "Seed hit (1 = yes)": bool01(a.seed_hit), "Status": a.status})).reset_index(drop=True)
    data_sheet(wb, "Attribution", "Appendix C", "Failure Bucket of Every Slice", out, fmt=fmt_for(out.columns),
               widths={"Measured value that set the bucket": 60, "Bucket meaning": 36})
    table_sheet(wb, "Tables", [tabs["buckets"]])
    return save(wb, d, "Appendix_C_E5_Failure_Attribution.xlsx")


def wb_e2(df, R, tabs, d):
    wb = Workbook()
    notes_sheet(wb, "Appendix D", "Experiment E2 — Objective 2: Minimizing Undersegmentation", [
        ("Objective", "To minimize undersegmentation by correcting the image using Log-Domain Transformation"),
        ("Configurations", "P_SRG, P_ESRG_global, P_ESRG_nolog, P_ESRG (same planted seed); "
                           "B_* = the same with the synthetic 40% bias field (Eq. 4.6)"),
        ("Primary metric", "Recall = TP / (TP + FN); Missed share = FN / |G| = 1 − Recall"),
        ("Sheets", "Recall_planted, Bias_field, Tables"),
    ] + metric_notes())
    b = base_cols(df)
    parts = [b]
    checks = []
    for cfg, pre in (("P_SRG", "SRG"), ("P_ESRG_global", "ESRG global"), ("P_ESRG_nolog", "ESRG log off"),
                     ("P_ESRG", "ESRG full")):
        parts.append(block(df, cfg, [("tp", "TP (px)"), ("fn", "FN (px)"), ("fp", "FP (px)"), ("recall", "Recall"),
                                     ("precision", "Precision"), ("dsc", "DSC")], pre))
        checks.append((f"{pre}: Recall (check)", f"={{{pre}: TP (px)}}/({{{pre}: TP (px)}}+{{{pre}: FN (px)}})"))
        checks.append((f"{pre}: Missed share (check)", f"={{{pre}: FN (px)}}/{{Ground-truth area |G| (px)}}"))
    out = parts[0].join(parts[1:]).reset_index(drop=True)
    data_sheet(wb, "Recall_planted", "Appendix D", "Recall per Slice With a Planted Seed", out,
               fmt=fmt_for(out.columns), checks=checks)
    parts = [b, block(df, "B_ESRG", [("bias_theta_deg", "θ (degrees)")], "Bias field")]
    checks = []
    for clean, biased, pre in (("P_SRG", "B_SRG", "SRG"), ("P_ESRG_global", "B_ESRG_global", "ESRG global"),
                               ("P_ESRG_nolog", "B_ESRG_nolog", "ESRG log off"), ("P_ESRG", "B_ESRG", "ESRG full")):
        parts.append(block(df, clean, [("recall", "Recall, no bias")], pre))
        parts.append(block(df, biased, [("tp", "TP with bias (px)"), ("fn", "FN with bias (px)"), ("recall", "Recall, bias field"),
                                        ("dsc", "DSC, bias field")], pre))
        checks.append((f"{pre}: change in recall (check)", f"={{{pre}: Recall, bias field}}-{{{pre}: Recall, no bias}}"))
    parts.append(block(df, "B_SRG_N4", [("recall", "Recall, bias field"), ("dsc", "DSC, bias field")], "SRG + N4"))
    out = parts[0].join(parts[1:]).reset_index(drop=True)
    fm = fmt_for(out.columns)
    fm["Bias field: θ (degrees)"] = F2
    data_sheet(wb, "Bias_field", "Appendix D", "Recall per Slice With and Without the Synthetic Bias Field", out,
               fmt=fm, checks=checks)
    table_sheet(wb, "Tables", [tabs["recall"], tabs["bias"]])
    return save(wb, d, "Appendix_D_E2_Undersegmentation.xlsx")


def wb_e3(df, R, tabs, d):
    wb = Workbook()
    notes_sheet(wb, "Appendix E", "Experiment E3 — Objective 3: Decreasing Boundary Leakage", [
        ("Objective", "To decrease boundary leakage by integrating an adaptive stopping criterion"),
        ("Configurations", "P_ESRG_nostop (unconditional absorption), P_SRG, P_ESRG (adaptive stopping); "
                           "P_ESRG_nopurify for the seed-size analysis"),
        ("Primary metrics", "Precision = TP/(TP + FP); Leaked = 1 when |M| > 2|G|; HD95 and ASSD in pixels"),
        ("Sheets", "Leakage_planted, Purification, Tables"),
    ] + metric_notes())
    b = base_cols(df)
    parts, checks = [b], []
    for cfg, pre in (("P_ESRG_nostop", "No stopping"), ("P_SRG", "SRG"), ("P_ESRG", "ESRG full")):
        parts.append(block(df, cfg, [("pred_area", "|M| (px)"), ("tp", "TP (px)"), ("fp", "FP (px)"), ("precision", "Precision"),
                                     ("area_ratio", "Area ratio |M|/|G|"), ("leaked", "Leaked (1 = yes)"),
                                     ("hd95_wc", "HD95 (px)"), ("assd_wc", "ASSD (px)"), ("stop_reason", "Stop reason")], pre))
        checks.append((f"{pre}: Precision (check)", f"={{{pre}: TP (px)}}/({{{pre}: TP (px)}}+{{{pre}: FP (px)}})"))
        checks.append((f"{pre}: Leaked (check)", f"=IF({{{pre}: |M| (px)}}>2*{{Ground-truth area |G| (px)}},1,0)"))
    out = parts[0].join(parts[1:]).reset_index(drop=True)
    for c in out.columns:
        if "Leaked" in c:
            out[c] = bool01(out[c])
    data_sheet(wb, "Leakage_planted", "Appendix E", "Precision, Leakage, and Boundary Distances per Slice (Planted Seed)",
               out, fmt=fmt_for(out.columns), checks=checks)
    parts = [b]
    for cfg, pre in (("P_ESRG_nopurify", "Click disk"), ("P_ESRG", "Purified core")):
        parts.append(block(df, cfg, [("seed_area", "seed area (px)"), ("sigma_A_initial", "initial σA"),
                                     ("sigma_floor", "σ floor"), ("dsc", "DSC"), ("recall", "Recall"),
                                     ("precision", "Precision"), ("leaked", "Leaked (1 = yes)")], pre))
    out = parts[0].join(parts[1:]).reset_index(drop=True)
    for c in out.columns:
        if "Leaked" in c:
            out[c] = bool01(out[c])
    fm = fmt_for(out.columns)
    for c in out.columns:
        if "σ" in c:
            fm[c] = "0.00000"
    data_sheet(wb, "Purification", "Appendix E", "Seed Size, Initial σA, and Outcome With and Without Seed Purification",
               out, fmt=fm)
    table_sheet(wb, "Tables", [tabs["leak"], tabs["boundary"], tabs["purify"]])
    return save(wb, d, "Appendix_E_E3_Boundary_Leakage.xlsx")


def wb_e4(df, R, tabs, d):
    wb = Workbook()
    notes_sheet(wb, "Appendix F", "Experiment E4 — Overall Delineation Quality and Ablation", [
        ("Configurations", "A_SRG, A_ESRG (automatic seed); P_SRG, P_ESRG (planted seed); ablation arms "
                           "P_ESRG_global, P_ESRG_nolog, P_ESRG_nostop"),
        ("Tests", "Wilcoxon signed-rank (Eq. 4.12), r = Z/√N (Eq. 3.41), Holm (Eq. 3.40), McNemar (Eq. 4.13), "
                  "Friedman (Eq. 4.16), Kruskal–Wallis, Shapiro–Wilk"),
        ("Population weighting (Eq. 4.15)", "ȳ_st = Σ_h W_h ȳ_h, W_h = N_h/N"),
        ("Sheets", "Overall_auto, Overall_planted, Ablation, Tables"),
    ] + metric_notes())
    b = base_cols(df)
    for key, pairs in (("Overall_auto", (("A_SRG", "SRG auto"), ("A_ESRG", "ESRG auto"))),
                       ("Overall_planted", (("P_SRG", "SRG planted"), ("P_ESRG", "ESRG planted")))):
        parts, checks = [b], []
        for cfg, pre in pairs:
            parts.append(block(df, cfg, OVERLAP, pre))
            checks += overlap_checks(pre)
        out = parts[0].join(parts[1:]).reset_index(drop=True)
        for c in out.columns:
            if "(1 = yes)" in c:
                out[c] = bool01(out[c])
        data_sheet(wb, key, "Appendix F", f"Overlap and Boundary Metrics per Slice ({'Automatic' if 'auto' in key else 'Planted'} Seed)",
                   out, fmt=fmt_for(out.columns), checks=checks)
    arms = [("P_ESRG", "ESRG full"), ("P_ESRG_global", "Global measure"), ("P_ESRG_nolog", "Log off"),
            ("P_ESRG_nostop", "No stopping")]
    m = pd.concat([df[df.config == c].set_index("file").dsc.rename(f"{n}: DSC") for c, n in arms], axis=1)
    rk = m.rank(axis=1, ascending=False)
    rk.columns = [c.replace("DSC", "rank (1 = best)") for c in m.columns]
    out = b.join(m).join(rk).reset_index(drop=True)
    data_sheet(wb, "Ablation", "Appendix F", "DSC and Within-Slice Rank of the Four ESRG Ablation Arms (Friedman Test)",
               out, fmt={**fmt_for(m.columns), **{c: F2 for c in rk.columns}})
    table_sheet(wb, "Tables", [tabs[k] for k in ("metrics", "normality", "overall_auto", "overall_planted",
                                                  "weighted", "dsc_class", "dsc_plane", "ablation")])
    return save(wb, d, "Appendix_F_E4_Overall_Delineation.xlsx")


def wb_e6(timing, R, tabs, d):
    if timing is None:
        return None
    wb = Workbook()
    env = R["meta"].get("environment") or {}
    notes_sheet(wb, "Appendix G", "Experiment E6 — Processing Time", [
        ("Configurations", "A_SRG and A_ESRG (automatic seeding)"),
        ("Slices", f"{timing.file.nunique()} slices: the first {timing.file.nunique() // 9} slices of every stratum in "
                   "the random draw order of the sampling step (a stratified random subsample of the 954)"),
        ("Protocol", "One process, nothing else running; one warm-up run discarded; order of SRG and ESRG "
                     "alternated per slice (Order column); time.perf_counter around the whole pipeline"),
        ("Excluded", "Computation of the evaluation metrics"),
    ] + [(k, str(v)) for k, v in env.items()])
    t = timing.copy()
    stages = [("t_input", "Input (s)"), ("t_mask", "Head mask (s)"), ("t_log", "Log domain (s)"),
              ("t_candidates", "Seed selection (s)"), ("t_growth", "Region growing (s)"), ("t_final", "Post-processing (s)")]
    parts = []
    for cfg, pre in (("A_SRG", "SRG"), ("A_ESRG", "ESRG")):
        s = t[t.config == cfg].set_index("file")
        parts.append(pd.DataFrame({f"{pre}: {n}": s[c] for c, n in [("seconds", "Total (s)")] + stages + [("order", "order (1 = first)")]}))
    s0 = t[t.config == "A_SRG"].set_index("file")
    base = pd.DataFrame({"Slice (file)": s0.index, "Split": s0.split.values, "Tumor class": s0.tumor.map(CLS).values,
                         "Plane": s0.plane.map(PLN).values}).set_index("Slice (file)", drop=False)
    out = base.join(parts[0]).join(parts[1]).sort_index().reset_index(drop=True)
    fm = {c: "0.0000" for c in out.columns if "(s)" in c}
    data_sheet(wb, "Time_per_slice", "Appendix G", "Sequential Processing Time per Slice and Stage", out, fmt=fm,
               checks=[("ESRG − SRG, s (check)", "={ESRG: Total (s)}-{SRG: Total (s)}"),
                       ("ESRG faster (check)", "=IF({ESRG: Total (s)}<{SRG: Total (s)},1,0)")])
    table_sheet(wb, "Tables", [tabs["time"], tabs["stages"]])
    return save(wb, d, "Appendix_G_E6_Processing_Time.xlsx")


COLUMN_DOC = {
    "file": "Slice filename", "split": "BRISC split the slice came from", "tumor": "Tumor class", "plane": "Imaging plane",
    "stratum": "Tumor class | plane", "index": "Slice number in the filename", "config": "Configuration id (Configurations sheet)",
    "method": "esrg or srg", "seeding": "auto, planted, or operator", "biased": "Synthetic bias field applied",
    "click": "Operator click number (1–5)", "status": "OK, WARN, NO TUMOR CANDIDATE, or ERROR",
    "tp": "True-positive pixels |M ∩ G|", "fp": "False-positive pixels |M \\ G|", "fn": "False-negative pixels |G \\ M|",
    "pred_area": "|M| (px)", "gt_area": "|G| (px)", "dsc": "Eq. 3.32", "iou": "Eq. 3.33", "precision": "Eq. 3.34",
    "recall": "Eq. 3.35", "hd95": "Eq. 3.37 (blank when a mask is empty)", "assd": "Eq. 4.10 (blank when a mask is empty)",
    "hd95_wc": "HD95 with the image diagonal for an empty prediction", "assd_wc": "ASSD with the image diagonal for an empty prediction",
    "diag": "Image diagonal (px)", "n_boundary_pred": "|∂M| (px)", "n_boundary_gt": "|∂G| (px)", "area_ratio": "|M|/|G|",
    "leaked": "|M| > 2|G| (Eq. 3.39)", "success": "DSC ≥ 0.70 (Eq. 4.11)", "seed_area": "|S| (px)", "seed_in_gt_px": "|S ∩ G| (px)",
    "seed_in_gt_frac": "|S ∩ G|/|S|", "seed_hit": "S ⊆ G (Eq. 3.38)", "seed_distance": "Eq. 4.7 (px)",
    "click_row": "Row of the click (planted/operator)", "click_col": "Column of the click", "click_depth": "Distance of the click to the tumor boundary (px)",
    "sigma_floor": "σ_floor (Eq. 3.21)", "sigma_A_initial": "σ_A in pass 1 (Eq. 3.28)", "k_local": "k_L", "stop_reason": "Why growth ended",
    "n_passes": "Growing passes", "bias_theta_deg": "Bias gradient direction (degrees)", "bucket": "E5 bucket (A_ESRG)",
    "bucket_reason": "Value that set the bucket", "seconds": "Time measured during the parallel run (not used for E6)",
    "t_input": "Stage 1 time (s)", "t_mask": "Stage 2 time (s)", "t_log": "Stage 3 time (s)", "t_candidates": "Stages 4–5 time (s)",
    "t_growth": "Stage 6 time (s)", "t_final": "Stage 7 time (s)", "error": "Exception text when status = ERROR",
}


def wb_raw(df, R, d):
    wb = Workbook()
    notes_sheet(wb, "Appendix H", "Raw Results of Every Run", [
        ("Rows", f"{len(df):,} runs = {R['meta']['n_sample']:,} slices × {R['meta']['n_configs']} configurations"),
        ("Errors", f"{R['meta']['n_errors']} runs ended in an error"),
        ("Reproduce", "python experiments/sampling.py --draw; python experiments/evaluate.py; "
                      "python experiments/evaluate.py --timing; python experiments/analyze.py"),
    ] + [(k, v) for k, v in COLUMN_DOC.items()])
    cfg = pd.DataFrame([{"Configuration": k, "Description": v[0], "Parameter overrides": str(v[1]) or "—",
                         "Seeding": v[2], "Bias field": v[3]} for k, v in CONFIGS.items()])
    data_sheet(wb, "Configurations", "Appendix H", "Configurations", cfg, widths={"Description": 50, "Parameter overrides": 40})
    raw = df.sort_values(["file", "config"]).copy()
    for c in ("leaked", "success", "seed_hit", "biased"):
        raw[c] = bool01(raw[c])
    data_sheet(wb, "Raw_results", "Appendix H", "One Row per Slice and Configuration", raw)
    return save(wb, d, "Appendix_H_Raw_Results.xlsx")


def write_all(df, timing, R, d):
    tabs = T.build(R)
    paths = [wb_sampling(R, tabs, d), wb_e1(df, R, tabs, d), wb_e5(df, R, tabs, d), wb_e2(df, R, tabs, d),
             wb_e3(df, R, tabs, d), wb_e4(df, R, tabs, d), wb_e6(timing, R, tabs, d), wb_raw(df, R, d)]
    for p in paths:
        if p:
            print("  wrote", p)
    return paths
