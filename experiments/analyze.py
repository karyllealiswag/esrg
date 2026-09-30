"""
analyze.py — Statistics, figures, and the raw-results workbook for Chapter 4.

Purpose : Turn the raw per-image rows written by evaluate.py into every number,
          table, and figure reported in Chapter 4, so that each value traces back
          to the raw results.
Function : Descriptive statistics per configuration; paired Wilcoxon signed-rank
          tests with Z, r = Z/sqrt(N), and Holm adjustment; exact McNemar tests for
          paired rates; Wilson intervals; chi-square, Mann-Whitney, Kruskal-Wallis,
          and Friedman tests. Writes results.json, PNG figures, and an .xlsx
          workbook holding the sample, every raw row, and a column dictionary.
Notes   : Pairs with a missing value (e.g. HD95 of an empty mask) are dropped per
          test. CLI: python experiments/analyze.py [--dir outputs/evaluation]
"""
import argparse
import json
import math
import os
import sys

import numpy as np
import pandas as pd
from scipy import stats

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from experiments.evaluate import CONFIGS, FIELDS, SAMPLE_FRAC, SAMPLE_SIZE, SAMPLE_SEED, INU
from experiments.attribution import BUCKET_LABEL

CLASSES = ["meningioma", "pituitary", "glioma"]
PLANES = ["axial", "coronal", "sagittal"]
METRICS = ["dsc", "iou", "precision", "recall", "hd95"]
LABEL = {"dsc": "DSC", "iou": "IoU", "precision": "Precision", "recall": "Recall", "hd95": "HD95 (px)"}


# ─── Statistics ──────────────────────────────────────────────────────────────
def desc(v):
    v = pd.to_numeric(pd.Series(v), errors="coerce").dropna().to_numpy(float)
    if v.size == 0:
        return {"n": 0}
    return {"n": int(v.size), "mean": float(v.mean()),
            "sd": float(v.std(ddof=1)) if v.size > 1 else 0.0,
            "median": float(np.median(v)), "q1": float(np.percentile(v, 25)),
            "q3": float(np.percentile(v, 75)),
            "ci95": float(stats.t.ppf(0.975, v.size - 1) * v.std(ddof=1) / math.sqrt(v.size)) if v.size > 1 else 0.0}


def wilcoxon(x, y):
    """
    Paired Wilcoxon signed-rank test of y against x (d = y - x). Zero differences
    are dropped (Wilcoxon's method); ranks of |d| use average ranks for ties.
    Z = (W+ - n'(n'+1)/4) / sqrt(n'(n'+1)(2n'+1)/24 - sum(t^3 - t)/48), two-sided
    p from the normal distribution, and r = Z / sqrt(N) with N the complete pairs.
    """
    x = pd.to_numeric(pd.Series(x), errors="coerce").to_numpy(float)
    y = pd.to_numeric(pd.Series(y), errors="coerce").to_numpy(float)
    ok = ~(np.isnan(x) | np.isnan(y))
    d = (y - x)[ok]
    N = int(ok.sum())
    nz = d[d != 0]
    n = nz.size
    out = {"N": N, "n_nonzero": int(n), "n_pos": int((nz > 0).sum()), "n_neg": int((nz < 0).sum())}
    if n == 0:
        return {**out, "W_plus": 0.0, "Z": 0.0, "p": 1.0, "r": 0.0}
    ranks = stats.rankdata(np.abs(nz))
    w_plus = float(ranks[nz > 0].sum())
    _, t = np.unique(np.abs(nz), return_counts=True)
    var = n * (n + 1) * (2 * n + 1) / 24 - float(((t ** 3) - t).sum()) / 48
    z = (w_plus - n * (n + 1) / 4) / math.sqrt(var) if var > 0 else 0.0
    p = float(2 * stats.norm.sf(abs(z)))
    return {**out, "W_plus": w_plus, "Z": float(z), "p": p, "r": float(z / math.sqrt(N))}


def holm(ps):
    """Holm step-down adjusted p-values (same order as input)."""
    m = len(ps)
    order = sorted(range(m), key=lambda i: ps[i])
    adj, run = [0.0] * m, 0.0
    for k, i in enumerate(order):
        run = max(run, min(1.0, (m - k) * ps[i]))
        adj[i] = run
    return adj


def mcnemar(a, b):
    """Exact McNemar test on paired booleans a (reference) and b (compared)."""
    a, b = np.asarray(a, bool), np.asarray(b, bool)
    n10, n01 = int((a & ~b).sum()), int((~a & b).sum())
    p = float(stats.binomtest(min(n10, n01), n10 + n01, 0.5).pvalue) if n10 + n01 else 1.0
    return {"only_ref": n10, "only_cmp": n01, "p": p}


def wilson(k, n, z=1.959964):
    if n == 0:
        return (float("nan"), float("nan"))
    p = k / n
    den = 1 + z * z / n
    c = (p + z * z / (2 * n)) / den
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return (100 * (c - h), 100 * (c + h))


def rate(v):
    v = [bool(x) for x in v if x is not None and not (isinstance(x, float) and math.isnan(x))]
    return 100.0 * sum(v) / len(v) if v else float("nan")


# ─── Data ────────────────────────────────────────────────────────────────────
def load(d):
    df = pd.read_csv(os.path.join(d, "raw_results.csv"))
    df = df[df.file.isin(set(pd.read_csv(os.path.join(d, "sample.csv")).file))]
    for c in ("leaked", "success", "seed_hit", "biased"):
        df[c] = df[c].map({True: True, False: False, "True": True, "False": False})
    return df


def wide(df, cfg, col):
    s = df[df.config == cfg].set_index("file")[col]
    return s


def paired(df, a, b, col, mask=None):
    x, y = wide(df, a, col), wide(df, b, col)
    idx = x.index.intersection(y.index)
    if mask is not None:
        idx = idx.intersection(mask)
    return x.loc[idx], y.loc[idx]


def compare(df, base, enh, rows_idx=None, holm_family=True):
    """Overall table: descriptives of both configs, mean difference, Wilcoxon, McNemar."""
    out, ps = [], []
    for m in METRICS:
        x, y = paired(df, base, enh, m, rows_idx)
        w = wilcoxon(x, y)
        out.append({"metric": m, "base": desc(x), "enh": desc(y),
                    "diff_mean": desc(y)["mean"] - desc(x)["mean"] if desc(x)["n"] else None, **w})
        ps.append(w["p"])
    for m in ("success", "leaked"):
        x, y = paired(df, base, enh, m, rows_idx)
        x, y = x.fillna(False).astype(bool), y.fillna(False).astype(bool)
        mc = mcnemar(x, y)
        out.append({"metric": m, "base_rate": 100 * x.mean(), "enh_rate": 100 * y.mean(),
                    "base_k": int(x.sum()), "enh_k": int(y.sum()), "N": int(len(x)),
                    "diff_pp": 100 * (y.mean() - x.mean()), **mc})
        ps.append(mc["p"])
    if holm_family:
        for row, pa in zip(out, holm(ps)):
            row["p_holm"] = pa
    return out


def per_class_tests(df, base, enh, col, classes=CLASSES):
    res, ps = {}, []
    for c in classes:
        idx = df[(df.config == enh) & (df.tumor == c)].file
        x, y = paired(df, base, enh, col, pd.Index(idx))
        w = wilcoxon(x, y)
        res[c] = {"base": desc(x), "enh": desc(y), **w}
        ps.append(w["p"])
    for c, pa in zip(classes, holm(ps)):
        res[c]["p_holm"] = pa
    x, y = paired(df, base, enh, col)
    res["all"] = {"base": desc(x), "enh": desc(y), **wilcoxon(x, y)}
    res["all"]["p_holm"] = res["all"]["p"]
    return res


def analyze(d):
    df = load(d)
    sample = pd.read_csv(os.path.join(d, "sample.csv"))
    strata = pd.read_csv(os.path.join(d, "sample_strata.csv"))
    R = {"meta": {"n_sample": int(len(sample)), "N_population": int(strata.N_h.sum()),
                  "frac": SAMPLE_FRAC, "size": SAMPLE_SIZE, "seed": SAMPLE_SEED, "inu": INU,
                  "n_rows": int(len(df)), "n_errors": int((df.status == "ERROR").sum()),
                  "errors": df[df.status == "ERROR"][["file", "config", "error"]].to_dict("records")}}
    R["strata"] = strata.to_dict("records")
    R["configs"] = [{"id": k, "desc": v[0], "seeding": v[2], "biased": v[3],
                     "overrides": v[1], "n": int((df.config == k).sum()),
                     "n_no_candidate": int(((df.config == k) & (df.status == "NO TUMOR CANDIDATE")).sum()),
                     "n_leak_warn": int(((df.config == k) & (df.status == "WARN")).sum())}
                    for k, v in CONFIGS.items()]

    # Summary of every configuration (E4 ablation overview)
    summ = {}
    for k in CONFIGS:
        s = df[df.config == k]
        summ[k] = {**{m: desc(s[m]) for m in METRICS},
                   "success": rate(s.success), "leaked": rate(s.leaked),
                   "area_ratio_median": float(s.area_ratio.median()),
                   "seed_area_median": float(s.seed_area.median()),
                   "sigma_A_initial_median": float(s.sigma_A_initial.median()) if s.sigma_A_initial.notna().any() else None,
                   "seconds_parallel": desc(s.seconds)}
    R["summary"] = summ

    # 4.1 overall comparisons
    R["overall_auto"] = compare(df, "A_SRG", "A_ESRG")
    R["overall_planted"] = compare(df, "P_SRG", "P_ESRG")
    # median area ratio (predicted / true area) per seeding
    R["area_ratio"] = {k: float(df[df.config == k].area_ratio.median()) for k in CONFIGS}

    # 4.2.1 Objective 1 — seed hit rate
    a = df[df.config == "A_ESRG"].copy()
    a["seeded"] = a.seed_area > 0
    obj1 = {}
    for grp, key in [("class", "tumor"), ("plane", "plane")]:
        levels = CLASSES if key == "tumor" else PLANES
        tab = {}
        for lv in levels + ["all"]:
            s = a if lv == "all" else a[a[key] == lv]
            seeded = s[s.seeded]
            k = int(seeded.seed_hit.fillna(False).astype(bool).sum())
            ns = int(len(seeded))
            hit = seeded[seeded.seed_hit == True]
            miss = s[~(s.seed_hit == True)]
            tab[lv] = {"n": int(len(s)), "N_S": ns, "no_candidate": int((~s.seeded).sum()),
                       "hits": k, "rate": 100 * k / ns if ns else float("nan"), "ci": wilson(k, ns),
                       "dsc_hit_median": float(hit.dsc.median()) if len(hit) else None,
                       "dsc_miss_median": float(miss.dsc.median()) if len(miss) else None,
                       "success_hit": rate(hit.success), "success_miss": rate(miss.success),
                       "seed_in_gt_mean": float(seeded.seed_in_gt_frac.mean())}
        cont = [[tab[lv]["hits"], tab[lv]["N_S"] - tab[lv]["hits"]] for lv in levels]
        try:
            chi2, p, dof, _ = stats.chi2_contingency(cont, correction=False)
            tab["chi2"] = {"chi2": float(chi2), "p": float(p), "dof": int(dof)}
        except ValueError:                       # a level with no seeded slices
            tab["chi2"] = None
        obj1[grp] = tab
    hit_d = a[a.seed_hit == True].dsc.dropna()
    miss_d = a[~(a.seed_hit == True)].dsc.dropna()
    mw = stats.mannwhitneyu(hit_d, miss_d, alternative="two-sided")
    obj1["mannwhitney"] = {"U": float(mw.statistic), "p": float(mw.pvalue),
                           "n_hit": int(hit_d.size), "n_miss": int(miss_d.size),
                           "r_rb": float(2 * mw.statistic / (hit_d.size * miss_d.size) - 1)}
    # automatic vs planted DSC per class for ESRG
    obj1["auto_vs_planted"] = per_class_tests(df, "P_ESRG", "A_ESRG", "dsc")
    R["obj1"] = obj1

    # E5 failure attribution
    buckets = {}
    for lv in CLASSES + ["all"]:
        s = a if lv == "all" else a[a.tumor == lv]
        vc = s.bucket.value_counts()
        buckets[lv] = {b: {"k": int(vc.get(b, 0)), "pct": 100 * vc.get(b, 0) / len(s)} for b in BUCKET_LABEL}
        buckets[lv]["n"] = int(len(s))
    R["buckets"] = buckets
    R["bucket_labels"] = BUCKET_LABEL

    # 4.2.2 Objective 2 — recall ablation (planted)
    obj2 = {}
    for cmp_cfg in ("P_SRG", "P_ESRG_global", "P_ESRG_nolog"):
        obj2[cmp_cfg] = {m: per_class_tests(df, cmp_cfg, "P_ESRG", m) for m in ("recall", "dsc", "precision")}
    # Holm across the three comparisons, per class, on recall
    for lv in CLASSES + ["all"]:
        ps = [obj2[c]["recall"][lv]["p"] for c in ("P_SRG", "P_ESRG_global", "P_ESRG_nolog")]
        for c, pa in zip(("P_SRG", "P_ESRG_global", "P_ESRG_nolog"), holm(ps)):
            obj2[c]["recall"][lv]["p_holm3"] = pa
    R["obj2"] = obj2

    # E2 bias robustness: clean vs biased, same planted seed
    e2, ps = {}, []
    for clean, biased in [("P_SRG", "B_SRG"), ("P_ESRG", "B_ESRG"),
                          ("P_ESRG_global", "B_ESRG_global"), ("P_ESRG_nolog", "B_ESRG_nolog")]:
        e2[biased] = {}
        for m in ("recall", "dsc", "precision"):
            x, y = paired(df, clean, biased, m)
            e2[biased][m] = {"clean": desc(x), "biased": desc(y), **wilcoxon(x, y),
                             "mean_abs_change": float(np.nanmean(np.abs(y.to_numpy(float) - x.to_numpy(float))))}
        e2[biased]["leaked"] = {"clean": rate(wide(df, clean, "leaked")), "biased": rate(wide(df, biased, "leaked"))}
        ps.append(e2[biased]["recall"]["p"])
    for k, pa in zip(e2, holm(ps)):
        e2[k]["recall"]["p_holm"] = pa
    x, y = paired(df, "B_SRG", "B_SRG_N4", "recall")
    e2["n4_vs_srg_biased"] = {"recall": {"srg": desc(x), "srg_n4": desc(y), **wilcoxon(x, y)},
                              "dsc": {"srg": desc(wide(df, "B_SRG", "dsc")), "srg_n4": desc(wide(df, "B_SRG_N4", "dsc"))}}
    x, y = paired(df, "B_SRG_N4", "B_ESRG", "recall")
    e2["esrg_vs_n4_biased"] = {"recall": wilcoxon(x, y)}
    R["e2"] = e2

    # 4.2.3 Objective 3 — stopping criterion (planted)
    obj3 = {"precision": per_class_tests(df, "P_ESRG_nostop", "P_ESRG", "precision"),
            "dsc": per_class_tests(df, "P_ESRG_nostop", "P_ESRG", "dsc"),
            "recall": per_class_tests(df, "P_ESRG_nostop", "P_ESRG", "recall"), "leak": {}}
    for lv in CLASSES + ["all"]:
        idx = df[(df.config == "P_ESRG") & ((df.tumor == lv) if lv != "all" else True)].file
        x, y = paired(df, "P_ESRG_nostop", "P_ESRG", "leaked", pd.Index(idx))
        x, y = x.fillna(False).astype(bool), y.fillna(False).astype(bool)
        xr, yr = paired(df, "P_ESRG_nostop", "P_ESRG", "area_ratio", pd.Index(idx))
        obj3["leak"][lv] = {"nostop": 100 * x.mean(), "esrg": 100 * y.mean(), **mcnemar(x, y),
                            "ratio_nostop": float(xr.median()), "ratio_esrg": float(yr.median()),
                            "n": int(len(x))}
    obj3["precision_median_esrg"] = float(wide(df, "P_ESRG", "precision").median())
    R["obj3"] = obj3

    # Seed purification trade-off (planted, ESRG)
    pur = compare(df, "P_ESRG_nopurify", "P_ESRG")
    R["purify"] = {"table": pur,
                   "seed_area": {"nopurify": float(wide(df, "P_ESRG_nopurify", "seed_area").median()),
                                 "purify": float(wide(df, "P_ESRG", "seed_area").median())},
                   "sigma_A": {"nopurify": desc(wide(df, "P_ESRG_nopurify", "sigma_A_initial")),
                               "purify": desc(wide(df, "P_ESRG", "sigma_A_initial"))},
                   "sigma_floor_binding": {
                       k: 100 * float((df[df.config == k].sigma_A_initial <= df[df.config == k].sigma_floor + 1e-9).mean())
                       for k in ("P_ESRG", "P_ESRG_nopurify", "A_ESRG")}}

    # 4.2.4 DSC by class and plane
    byc = {}
    for lv in CLASSES + ["all"]:
        row = {}
        for k in ("A_SRG", "A_ESRG", "P_SRG", "P_ESRG"):
            s = df[(df.config == k) & ((df.tumor == lv) if lv != "all" else True)]
            row[k] = {"median": float(s.dsc.median()), "mean": float(s.dsc.mean()),
                      "success": rate(s.success), "hd95_median": float(s.hd95.median()),
                      "dsc0": 100 * float((s.dsc == 0).mean()), "n": int(len(s))}
        byc[lv] = row
    R["dsc_by_class"] = byc
    R["dsc_tests_auto"] = per_class_tests(df, "A_SRG", "A_ESRG", "dsc")
    R["dsc_tests_planted"] = per_class_tests(df, "P_SRG", "P_ESRG", "dsc")
    R["hd95_tests_auto"] = per_class_tests(df, "A_SRG", "A_ESRG", "hd95")
    kw = {}
    for k in ("A_ESRG", "P_ESRG"):
        s = df[df.config == k]
        kw[k] = {}
        for name, col, levels in (("class", "tumor", CLASSES), ("plane", "plane", PLANES)):
            groups = [g for g in (s[s[col] == lv].dsc.dropna() for lv in levels) if len(g)]
            h = stats.kruskal(*groups) if len(groups) > 1 else None
            kw[k][name] = {"H": float(h.statistic), "p": float(h.pvalue)} if h else None
    R["kruskal"] = kw
    byp = {}
    for c in CLASSES + ["all"]:
        for p in PLANES + ["all"]:
            row = {}
            for k in ("A_ESRG", "P_ESRG"):
                s = df[(df.config == k) & ((df.tumor == c) if c != "all" else True)
                       & ((df.plane == p) if p != "all" else True)]
                row[k] = {"median": float(s.dsc.median()), "success": rate(s.success), "n": int(len(s))}
            byp[f"{c}|{p}"] = row
    R["dsc_by_class_plane"] = byp

    # E4 Friedman across the planted ESRG ablation arms (DSC)
    arms = ["P_ESRG", "P_ESRG_global", "P_ESRG_nolog", "P_ESRG_nostop"]
    mat = pd.concat([wide(df, k, "dsc").rename(k) for k in arms], axis=1).dropna()
    fr = stats.friedmanchisquare(*[mat[k] for k in arms])
    R["friedman"] = {"arms": arms, "chi2": float(fr.statistic), "p": float(fr.pvalue), "n": int(len(mat)),
                     "mean_rank": {k: float(v) for k, v in mat.rank(axis=1, ascending=False).mean().items()}}

    # Pairwise follow-up to Friedman: full ESRG vs each ablated arm (DSC), Holm over 3
    pw = {k: wilcoxon(mat[k], mat["P_ESRG"]) for k in arms[1:]}
    for k, pa in zip(pw, holm([pw[k]["p"] for k in pw])):
        pw[k]["p_holm"] = pa
    R["friedman"]["pairwise"] = pw

    # Supporting facts cited in the discussion
    ae = df[df.config == "A_ESRG"]
    leak = ae[ae.leaked == True]
    gt = df[df.config == "P_ESRG"].gt_area
    diam = 2 * np.sqrt(gt / np.pi)
    R["extra"] = {
        "auto_leak_n": int(len(leak)), "auto_leak_seed_miss": int((leak.seed_hit != True).sum()),
        "auto_dsc0_srg": 100 * float((df[df.config == "A_SRG"].dsc == 0).mean()),
        "auto_dsc0_esrg": 100 * float((ae.dsc == 0).mean()),
        "gt_area_median": float(gt.median()), "gt_diam_median": float(diam.median()),
        # A linear field changes by INU over the image diagonal (~sqrt(2) * 512 px at most)
        "bias_change_over_median_tumor_pct": float(100 * INU * diam.median() / (math.sqrt(2) * 512)),
        "stop_reasons_P_ESRG": df[df.config == "P_ESRG"].stop_reason.str.replace(r"[0-9.]+", "#", regex=True)
                               .value_counts().to_dict(),
    }

    # E6 processing time — sequential pass
    tp = os.path.join(d, "timing_sequential.csv")
    t = pd.read_csv(tp) if os.path.isfile(tp) else df[df.config.isin(["A_SRG", "A_ESRG"])]
    R["timing_source"] = "sequential" if os.path.isfile(tp) else "parallel run (10 worker processes)"
    if True:
        tt = {}
        for lv in CLASSES + ["all"]:
            s = t if lv == "all" else t[t.tumor == lv]
            xs = s[s.config == "A_SRG"].set_index("file").seconds
            ys = s[s.config == "A_ESRG"].set_index("file").seconds
            idx = xs.index.intersection(ys.index)
            tt[lv] = {"srg": desc(xs.loc[idx]), "esrg": desc(ys.loc[idx]), **wilcoxon(xs.loc[idx], ys.loc[idx])}
        R["timing"] = tt
    return df, R


# ─── Figures ─────────────────────────────────────────────────────────────────
C_ESRG, C_SRG, C_A3, C_A4, C_A5 = "#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4"
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"


def _style(plt):
    plt.rcParams.update({"font.family": "Times New Roman", "font.size": 10, "axes.edgecolor": INK2,
                         "axes.labelcolor": INK, "xtick.color": INK2, "ytick.color": INK2,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "axes.grid": True, "axes.grid.axis": "y", "grid.color": GRID,
                         "grid.linewidth": 0.8, "axes.axisbelow": True, "legend.frameon": False})


def _bars(ax, groups, series, values, errs, colors, ylim=(0, 1), fmt="{:.2f}"):
    n = len(series)
    w = 0.8 / n
    x = np.arange(len(groups))
    for i, (s, col) in enumerate(zip(series, colors)):
        pos = x - 0.4 + w * (i + 0.5)
        ax.bar(pos, values[i], w * 0.92, color=col, label=s, yerr=errs[i] if errs else None,
               error_kw={"elinewidth": 0.9, "capsize": 2.5, "ecolor": INK2})
        for p, v, e in zip(pos, values[i], errs[i] if errs else [0] * len(pos)):
            ax.text(p, v + (e or 0) + 0.015 * (ylim[1] - ylim[0]), fmt.format(v), ha="center",
                    va="bottom", fontsize=7, color=INK2)
    ax.set_xticks(x, groups)
    ax.set_ylim(*ylim)


def figures(df, R, outdir):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    _style(plt)
    os.makedirs(outdir, exist_ok=True)
    S = R["summary"]
    mets = ["dsc", "iou", "precision", "recall"]

    # Figure 4.1 — overall means, auto and planted
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.1), sharey=True)
    for ax, (b, e, title) in zip(axes, [("A_SRG", "A_ESRG", "(a) Automatic seeding"),
                                        ("P_SRG", "P_ESRG", "(b) Seed planted inside the tumor")]):
        vals = [[S[b][m]["mean"] for m in mets], [S[e][m]["mean"] for m in mets]]
        errs = [[S[b][m]["ci95"] for m in mets], [S[e][m]["ci95"] for m in mets]]
        _bars(ax, [LABEL[m] for m in mets], ["SRG (baseline)", "ESRG (enhanced)"], vals, errs, [C_SRG, C_ESRG], (0, 1.12))
        ax.set_yticks(np.arange(0, 1.01, 0.2))
        ax.set_title(title, fontsize=10, loc="left", color=INK)
    axes[0].set_ylabel("Mean score (95% CI)")
    h, l = axes[0].get_legend_handles_labels()
    fig.legend(h, l, loc="lower center", ncol=2, fontsize=9)
    fig.tight_layout(rect=(0, 0.08, 1, 1))
    fig.savefig(os.path.join(outdir, "fig4_1_overall.png"), dpi=300)
    plt.close(fig)

    # Figure 4.2 — seed hit rate by class, Wilson CI
    t = R["obj1"]["class"]
    lv = CLASSES + ["all"]
    vals = [t[k]["rate"] for k in lv]
    lo = [t[k]["rate"] - t[k]["ci"][0] for k in lv]
    hi = [t[k]["ci"][1] - t[k]["rate"] for k in lv]
    fig, ax = plt.subplots(figsize=(5.0, 3.0))
    x = np.arange(len(lv))
    ax.bar(x, vals, 0.55, color=C_ESRG, yerr=[lo, hi], error_kw={"elinewidth": 0.9, "capsize": 3, "ecolor": INK2})
    for i, v in enumerate(vals):
        ax.text(i, v + hi[i] + 1.5, f"{v:.1f}%", ha="center", fontsize=8, color=INK2)
    ax.set_xticks(x, ["Meningioma", "Pituitary", "Glioma", "All"])
    ax.set_ylim(0, 100)
    ax.set_ylabel("Seed hit rate, % (95% CI)")
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, "fig4_2_seed_hit.png"), dpi=300)
    plt.close(fig)

    # Figure 4.3 — Objective 2 recall ablation and Objective 3 precision
    def cls_means(cfg, m):
        out, err = [], []
        for c in CLASSES + ["all"]:
            s = df[(df.config == cfg) & ((df.tumor == c) if c != "all" else True)][m]
            dd = desc(s)
            out.append(dd["mean"]); err.append(dd["ci95"])
        return out, err
    fig, axes = plt.subplots(2, 1, figsize=(7.2, 6.4))
    groups = ["Meningioma", "Pituitary", "Glioma", "All"]
    cfgs = [("P_SRG", "SRG", C_SRG), ("P_ESRG_global", "ESRG, global measure", C_A3),
            ("P_ESRG_nolog", "ESRG, log transform off", C_A4), ("P_ESRG", "ESRG (full)", C_ESRG)]
    v, e = zip(*[cls_means(c, "recall") for c, _, _ in cfgs])
    _bars(axes[0], groups, [n for _, n, _ in cfgs], v, e, [c for _, _, c in cfgs], (0, 1.15))
    axes[0].set_title("(a) Objective 2: recall", fontsize=10, loc="left")
    axes[0].set_ylabel("Mean recall (95% CI)")
    axes[0].legend(loc="upper left", ncol=2, fontsize=8)
    cfgs3 = [("P_ESRG_nostop", "ESRG, no stopping criterion", C_A5), ("P_ESRG", "ESRG (full)", C_ESRG)]
    v, e = zip(*[cls_means(c, "precision") for c, _, _ in cfgs3])
    _bars(axes[1], groups, [n for _, n, _ in cfgs3], v, e, [c for _, _, c in cfgs3], (0, 1.3))
    axes[1].set_title("(b) Objective 3: precision", fontsize=10, loc="left")
    axes[1].set_ylabel("Mean precision (95% CI)")
    axes[1].legend(loc="upper left", ncol=2, fontsize=8)
    for ax in axes:
        ax.set_yticks(np.arange(0, 1.01, 0.2))
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, "fig4_3_obj2_obj3.png"), dpi=300)
    plt.close(fig)

    # Figure 4.4 — bias robustness: recall clean vs biased
    pairs = [("P_SRG", "B_SRG", "SRG"), ("P_ESRG_global", "B_ESRG_global", "ESRG,\nglobal measure"),
             ("P_ESRG_nolog", "B_ESRG_nolog", "ESRG,\nlog off"), ("P_ESRG", "B_ESRG", "ESRG (full)")]
    fig, ax = plt.subplots(figsize=(6.2, 3.1))
    clean = [S[a]["recall"]["mean"] for a, _, _ in pairs]
    biased = [S[b]["recall"]["mean"] for _, b, _ in pairs]
    ce = [S[a]["recall"]["ci95"] for a, _, _ in pairs]
    be = [S[b]["recall"]["ci95"] for _, b, _ in pairs]
    _bars(ax, [n for _, _, n in pairs], ["Without bias field", f"With {int(INU*100)}% bias field"],
          [clean, biased], [ce, be], ["#9ec5ef", C_ESRG], (0, 0.8))
    ax.set_ylabel("Mean recall (95% CI)")
    ax.legend(loc="upper left", fontsize=8)
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, "fig4_4_bias.png"), dpi=300)
    plt.close(fig)

    # Figure 4.5 — DSC distributions (bimodality)
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.9), sharey=True)
    bins = np.linspace(0, 1, 21)
    for ax, (b, e, title) in zip(axes, [("A_SRG", "A_ESRG", "(a) Automatic seeding"),
                                        ("P_SRG", "P_ESRG", "(b) Seed planted inside the tumor")]):
        for cfg, col, name in [(b, C_SRG, "SRG (baseline)"), (e, C_ESRG, "ESRG (enhanced)")]:
            v = df[df.config == cfg].dsc.dropna()
            ax.hist(v, bins=bins, weights=np.full(len(v), 100 / len(v)), histtype="step",
                    linewidth=2, color=col, label=name)
        ax.set_title(title, fontsize=10, loc="left")
        ax.set_xlabel("DSC")
    axes[0].set_ylabel("Share of slices (%)")
    axes[0].legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, "fig4_5_dsc_distribution.png"), dpi=300)
    plt.close(fig)


# ─── Workbook ────────────────────────────────────────────────────────────────
COLUMN_DOC = {
    "file": "Slice filename in the BRISC 2025 test split",
    "tumor": "Tumor class parsed from the filename (glioma, meningioma, pituitary)",
    "plane": "Imaging plane parsed from the filename (axial, coronal, sagittal)",
    "index": "Slice number in the BRISC filename",
    "config": "Configuration id (see the Configurations sheet)",
    "method": "Region-growing algorithm: esrg (enhanced) or srg (Adams & Bischof baseline)",
    "seeding": "auto = automated seed selection; planted = click at the deepest ground-truth pixel",
    "biased": "True when a synthetic 40% linear bias field was multiplied into the slice (E2)",
    "status": "Pipeline status: OK, WARN (mask > 40% of head), NO TUMOR CANDIDATE, or ERROR",
    "dsc": "Dice Similarity Coefficient, 2TP/(2TP+FP+FN) (Eq. 3.32)",
    "iou": "Intersection over Union, TP/(TP+FP+FN) (Eq. 3.33)",
    "precision": "TP/(TP+FP) (Eq. 3.34); empty when nothing was predicted",
    "recall": "TP/(TP+FN) (Eq. 3.35)",
    "hd95": "95th-percentile Hausdorff distance in pixels (Eq. 3.37); empty when a mask is empty",
    "assd": "Average symmetric surface distance in pixels (recorded, not reported)",
    "tp": "True-positive pixels", "fp": "False-positive pixels", "fn": "False-negative pixels",
    "pred_area": "Predicted tumor area |M| in pixels", "gt_area": "Ground-truth tumor area |G| in pixels",
    "area_ratio": "|M| / |G|",
    "leaked": "True when |M| > 2|G| (Eq. 3.39, lambda = 2)",
    "success": "True when DSC >= 0.70",
    "seed_hit": "True when the whole seed core lies inside the ground truth (Eq. 3.38); empty if no seed",
    "seed_area": "Seed core area in pixels after purification (ESRG) or as planted (SRG)",
    "seed_in_gt_frac": "Share of the seed core inside the ground truth",
    "planted_row": "Row of the planted click (deepest ground-truth pixel), planted seeding only",
    "planted_col": "Column of the planted click, planted seeding only",
    "sigma_floor": "Noise floor sigma_floor of the slice (Eq. 3.21)",
    "sigma_A_initial": "sigma_A used in pass 1, max(s_A of the seed, sigma_floor) (Eq. 3.28); ESRG only",
    "k_local": "k_L used by the run", "stop_reason": "Why region growing ended",
    "n_passes": "Number of growing passes executed (ESRG)",
    "bias_theta_deg": "Direction of the synthetic bias gradient in degrees (E2 only)",
    "bucket": "Failure-attribution bucket (E5), A_ESRG only",
    "bucket_reason": "Measured value that placed the slice in its bucket",
    "seconds": "Processing time in seconds, measured during the parallel run (see Timing sheet for E6)",
    "error": "Exception text when status = ERROR",
}


def workbook(df, R, d, path):
    from openpyxl.styles import Font, Alignment, PatternFill
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        readme = pd.DataFrame({"Item": [
            "Study", "Data", "Sample", "Random seed", "Configurations", "Rows",
            "Errors", "How to read", "Reproduce"], "Value": [
            "Enhanced Seeded Region Growing (ESRG) vs. Adams & Bischof (1994) SRG — Chapter 4 raw results",
            "BRISC 2025 segmentation task, test split (860 slices, all with a tumor mask)",
            f"Proportional stratified random sample of {R['meta']['n_sample']} slices "
            f"(n_h = 500 x N_h / 860 per tumor class x plane stratum; Strata sheet)",
            str(SAMPLE_SEED),
            f"{len(CONFIGS)} configurations per slice (Configurations sheet)",
            f"{R['meta']['n_rows']} rows in Raw_results (one per slice x configuration)",
            f"{R['meta']['n_errors']} rows with status ERROR",
            "Every Chapter 4 value is computed from Raw_results (and Timing_sequential for E6) by experiments/analyze.py",
            "python experiments/evaluate.py; python experiments/evaluate.py --timing; python experiments/analyze.py"]})
        readme.to_excel(xw, sheet_name="README", index=False)
        pd.DataFrame([{"config": c["id"], "description": c["desc"], "seeding": c["seeding"],
                       "bias field": c["biased"], "parameter overrides": json.dumps(c["overrides"]),
                       "rows": c["n"]} for c in R["configs"]]).to_excel(xw, sheet_name="Configurations", index=False)
        pd.DataFrame({"column": list(COLUMN_DOC), "meaning": list(COLUMN_DOC.values())}).to_excel(xw, sheet_name="Column_dictionary", index=False)
        pd.read_csv(os.path.join(d, "sample_strata.csv")).to_excel(xw, sheet_name="Strata", index=False)
        pd.read_csv(os.path.join(d, "sample.csv")).to_excel(xw, sheet_name="Sample", index=False)
        df.sort_values(["file", "config"])[[c for c in FIELDS if c in df]].to_excel(xw, sheet_name="Raw_results", index=False)
        tp = os.path.join(d, "timing_sequential.csv")
        if os.path.isfile(tp):
            pd.read_csv(tp).sort_values(["file", "config"]).to_excel(xw, sheet_name="Timing_sequential", index=False)
        rows = []
        for k, s in R["summary"].items():
            rows.append({"config": k, **{f"{m}_mean": s[m].get("mean") for m in METRICS},
                         **{f"{m}_median": s[m].get("median") for m in METRICS},
                         "success_%": s["success"], "leakage_%": s["leaked"],
                         "area_ratio_median": s["area_ratio_median"]})
        pd.DataFrame(rows).to_excel(xw, sheet_name="Summary_by_config", index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E7E6E6")
            for col in ws.columns:
                width = min(60, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2)
                ws.column_dimensions[col[0].column_letter].width = max(10, width)
            ws.freeze_panes = "A2"
        for ws in (xw.book["README"], xw.book["Column_dictionary"], xw.book["Configurations"]):
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")


def _clean(o):
    if isinstance(o, dict):
        return {str(k): _clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_clean(v) for v in o]
    if isinstance(o, (np.floating, float)):
        return None if math.isnan(float(o)) else float(o)
    if isinstance(o, (np.integer,)):
        return int(o)
    if isinstance(o, np.bool_):
        return bool(o)
    return o


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default="outputs/evaluation")
    args = ap.parse_args()
    df, R = analyze(args.dir)
    with open(os.path.join(args.dir, "results.json"), "w") as f:
        json.dump(_clean(R), f, indent=1)
    figures(df, R, os.path.join(args.dir, "figures"))
    workbook(df, R, args.dir, os.path.join(args.dir, "ESRG_Chapter4_raw_evaluation_results.xlsx"))
    print("results.json, figures/, and the workbook written to", args.dir)
